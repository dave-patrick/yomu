"""Generate an OOXML DOCX fixture with python-docx, including repeated headers/footers."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "results" / "docx-fixture.docx"


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.header.paragraphs[0].text = "REPEATED HEADER SHOULD NOT BE READ"
    section.footer.paragraphs[0].text = "REPEATED FOOTER SHOULD NOT BE READ"
    document.add_heading("DOCX reading order fixture", level=1)
    document.add_paragraph("The first paragraph should be read before the table.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Table cell alpha"
    table.cell(0, 1).text = "Table cell beta"
    document.add_paragraph("The final paragraph should be read last.")
    document.save(OUTPUT)
    print(f"{OUTPUT}\t{OUTPUT.stat().st_size}")


if __name__ == "__main__":
    main()
